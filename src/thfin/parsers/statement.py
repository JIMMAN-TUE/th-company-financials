"""Extract revenue / cost / profit from financial-statement tables.

The parser is deliberately generic: it reads tables as rows of cells, finds a header row that
contains fiscal years, then matches row labels (Thai or English) against FIELD_PATTERNS.
It works on HTML tables and Excel sheets. If a source labels things differently, extend
FIELD_PATTERNS (and add a test) rather than writing a new parser.
"""

from __future__ import annotations

import re
from pathlib import Path

from bs4 import BeautifulSoup

from ..models import FINANCIAL_FIELDS
from .numbers import parse_number

FIELD_PATTERNS: dict[str, re.Pattern[str]] = {
    "revenue_total": re.compile(r"รายได้รวม|รวมรายได้|total\s+revenues?|total\s+income", re.I),
    "revenue_sales": re.compile(
        r"รายได้หลัก|รายได้จากการขาย|รายได้จากการให้บริการ|sales\s+(and|&)\s+service|"
        r"revenues?\s+from\s+(sales|services?)|^\s*sales\s*$",
        re.I,
    ),
    "cost_of_sales": re.compile(
        r"ต้นทุนขาย|ต้นทุนการขาย|ต้นทุนการให้บริการ|cost\s+of\s+(sales|goods\s+sold|services?)", re.I
    ),
    "total_expenses": re.compile(r"รายจ่ายรวม|ค่าใช้จ่ายรวม|รวมค่าใช้จ่าย|total\s+expenses?", re.I),
    "net_profit": re.compile(
        r"กำไร\s*\(?ขาดทุน\)?\s*สุทธิ|กำไรสุทธิ|ขาดทุนสุทธิ|net\s+(profit|income|loss)", re.I
    ),
    "total_assets": re.compile(r"สินทรัพย์รวม|รวมสินทรัพย์|total\s+assets", re.I),
}

# "2567", "ปี 2567", "FY2024", "31/12/2567" -> year token. Money like "2,567" does not match.
_YEAR_CELL = re.compile(r"^(?:[^\d]*|\d{1,2}[/\-.]\d{1,2}[/\-.])\s*(25\d{2}|20\d{2})$")

Rows = list[list[str]]
YearData = dict[int, dict[str, float | None]]


def _to_be(year: int) -> int:
    return year + 543 if year < 2400 else year


def _empty_bucket() -> dict[str, float | None]:
    return {f: None for f in FINANCIAL_FIELDS}


def _header_years(cells: list[str]) -> dict[int, int]:
    """{column_index: fiscal_year_BE} if the row is a year header (years + text only), else {}."""
    years: dict[int, int] = {}
    for idx, cell in enumerate(cells):
        m = _YEAR_CELL.match(cell.strip())
        if m:
            years[idx] = _to_be(int(m.group(1)))
    if not years:
        return {}
    for idx, cell in enumerate(cells):
        if idx not in years and cell.strip() and parse_number(cell) is not None:
            return {}  # contains other numbers -> a data row, not a header
    return years


def extract_from_rows(rows: Rows) -> YearData:
    """Core extraction. The first value found per (year, field) wins."""
    result: YearData = {}
    header: dict[int, int] = {}
    for raw in rows:
        cells = [c.strip() for c in raw]
        if not any(cells):
            continue
        years = _header_years(cells)
        if years:
            header = years
            continue
        if not header:
            continue
        label = next((c for c in cells if c and parse_number(c) is None), "")
        if not label:
            continue
        for fld, pattern in FIELD_PATTERNS.items():
            if not pattern.search(label):
                continue
            for col, year in header.items():
                value = parse_number(cells[col]) if col < len(cells) else None
                bucket = result.setdefault(year, _empty_bucket())
                if bucket[fld] is None and value is not None:
                    bucket[fld] = value
            break  # a row feeds at most one field
    return {y: d for y, d in result.items() if any(v is not None for v in d.values())}


def _merge(into: YearData, new: YearData) -> None:
    for year, fields in new.items():
        bucket = into.setdefault(year, _empty_bucket())
        for fld, value in fields.items():
            if bucket.get(fld) is None and value is not None:
                bucket[fld] = value


def html_tables_to_rows(html: str) -> list[Rows]:
    """Convert every <table> into rows of text cells, expanding colspan to keep columns aligned."""
    soup = BeautifulSoup(html, "html.parser")
    tables: list[Rows] = []
    for table in soup.find_all("table"):
        rows: Rows = []
        for tr in table.find_all("tr"):
            cells: list[str] = []
            for td in tr.find_all(["td", "th"]):
                text = re.sub(r"\s+", " ", td.get_text(" ", strip=True))
                raw_span = str(td.get("colspan", "1"))
                span = min(int(raw_span), 20) if raw_span.isdigit() and int(raw_span) > 0 else 1
                cells.extend([text] + [""] * (span - 1))
            if cells:
                rows.append(cells)
        if rows:
            tables.append(rows)
    return tables


def parse_html(html: str) -> YearData:
    out: YearData = {}
    for rows in html_tables_to_rows(html):
        _merge(out, extract_from_rows(rows))
    return out


def parse_excel(path: str | Path) -> YearData:
    from openpyxl import load_workbook

    out: YearData = {}
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows: Rows = [
                ["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)
            ]
            _merge(out, extract_from_rows(rows))
    finally:
        wb.close()
    return out
